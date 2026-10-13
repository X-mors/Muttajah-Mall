# -*- coding: utf-8 -*-
"""
حصر خرسانة الأعمدة (Column CONCRETE takeoff).

يبني على columns_takeoff:
  • العدد وموقع كل عمود من لوحة COLUMNS & AXES.
  • مقطع كل ماركة لكل مرحلة من المراحل الخمس من جدول الأعمدة (المقطع يصغر صعوداً).
ثم لكل عمود:  الحجم = الطول × العرض × الارتفاع الصافي  (م³).

قاعدة الارتفاع (المهندس): من سقف الدور الأسفل حتى سقف الدور (سقف↔سقف)،
ناقص سُمك ما يعلو العمود: بلاطة 22سم، أو drop panel 32سم حيث يوجد.
المناسيب (S.S.L) قد تختلف من عمود لآخر داخل الدور الواحد (بلاطة مُدرَّجة)، فالارتفاع
يُؤخذ لكل عمود من أقرب منسوب له — يُدخلها المهندس أو تُستخرج لاحقاً.

المُخرَجات:
  • لوحة مراجعة (DWG) فيها دائرة ورقم تسلسلي على كل عمود.
  • شيت Excel: حصر كل عمود بالمتر المكعب + الإجمالي النهائي.

الاستخدام:
    python concrete_columns_takeoff.py --debug                 # مقاطع المراحل + استخراج آلي
    python concrete_columns_takeoff.py --stage 2 --height 4.60 --dwg   # حصر أعمدة مرحلة 2
"""
import argparse, os, re, math, collections
import ezdxf
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import columns_takeoff as ct           # collect, count_marks, mark_key, text_of, MARK_RE, CORE_RE, SIZE_RE
try:
    import takeoff                     # to_dwg (DXF→DWG عبر ODA)
except Exception:
    takeoff = None

DEFAULT_DXF = ct.DEFAULT_DXF
DEFAULT_OUT = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "output", "حصر خرسانة الأعمدة")
STAGE_HDR_RE = re.compile(r"^FROM\s+.+\b(CEILING|ROOF|FOUNDATION)", re.I)

# تعريب أسماء السقوف + ربط تسمية الجدول بتسمية لوحات الأسقف (للمناسيب)
GLOSS = [
    ("LOWER BASEMENT CEILING", "سقف البدروم السفلي"),
    ("UPPER BASEMENT CEILING", "سقف البدروم العلوي"),
    ("5TH CEILING", "سقف المتكرر (5TH)"),
    ("3RD CEILING", "سقف المتكرر (3RD)"),
    ("1ST CEILING", "سقف الأرضي"),
    ("TOP ROOF", "الروف العلوي"),
    ("ROOF", "الروف"),
    ("FOUNDATION", "الأساسات"),
    ("GROUND CEILING", "سقف الأرضي"),
]
SCHED_TO_PLAN = {"1ST CEILING": "GROUND CEILING", "3RD CEILING": "TYPICAL CEILING",
                 "5TH CEILING": "TYPICAL CEILING"}

def gloss(name):
    u = (name or "").upper()
    for k, v in GLOSS:
        if k in u: return v
    return name

# =================== قراءة جدول الأعمدة (المراحل الخمس) ===================
def stage_anchors(sched):
    """مراكز X لعناوين المراحل الخمس 'FROM ... TO ... CEILING' (يسار→يمين = مرحلة 1→5)."""
    hs = []
    for (x, y, t) in sched:
        tt = " ".join(t.split())
        if STAGE_HDR_RE.match(tt) and "SCHEDULE" not in tt.upper():
            hs.append((x, tt))
    uniq = {}
    for x, t in hs:
        uniq[round(x / 500) * 500] = (x, t)
    return sorted(uniq.values(), key=lambda p: p[0])      # [(x, name)] مرحلة 1..5

def parse_schedule_stages(sched):
    """لكل ماركة: قائمة بطول 5 = مقطع كل مرحلة (W,H) مم أو None (---- = منتهٍ)."""
    anchors = stage_anchors(sched)
    ax = [a[0] for a in anchors]; names = [a[1] for a in anchors]
    if not ax:
        return {}, []
    heads = []
    for (x, y, t) in sched:
        m = re.match(r"^C\s*\.?\s*(\d{0,3})$", t, re.I)
        if m and x < min(ax) - 300:
            heads.append((x, y, "C" + m.group(1)))
    out = {}
    for (hx, hy, mark) in heads:
        row = [(x, t) for (x, yy, t) in sched if hy - 1450 < yy < hy + 250]
        stages = [None] * len(ax)
        for (x, t) in row:
            ms = ct.SIZE_RE.search(t)
            if not ms: continue
            k = min(range(len(ax)), key=lambda i: abs(ax[i] - x))
            stages[k] = (int(ms.group(1)), int(ms.group(2)))
        out[mark] = stages
    return out, names

def section_area_m2(wh):
    if not wh: return 0.0
    return (wh[0] / 1000.0) * (wh[1] / 1000.0)

# =================== مواقع الأعمدة وترقيمها ===================
def columns_list(plan):
    """كل عمود فِعلي بموقعه؛ يُرتَّب أعلى→أسفل ثم يسار→يمين ويُرقَّم تسلسلياً."""
    cols = []
    for (x, y, t) in plan:
        mm = ct.MARK_RE.match(t); cm = ct.CORE_RE.match(t)
        if mm: cols.append(("C" + mm.group(1), x, y))
        elif cm: cols.append(("CORE" + cm.group(1), x, y))
    cols.sort(key=lambda c: (-round(c[2] / 1000), round(c[1] / 1000)))
    return [(i + 1, m, x, y) for i, (m, x, y) in enumerate(cols)]

def build_review(src_dxf, cols, out_dxf):
    """ينسخ اللوحة ويضع دائرة ورقماً تسلسلياً على كل عمود."""
    doc = ezdxf.readfile(src_dxf); msp = doc.modelspace()
    LAY = "CHECK-CONCRETE-COLUMNS"
    if LAY not in doc.layers: doc.layers.add(LAY, color=3)
    R = 700
    for (sn, m, x, y) in cols:
        msp.add_circle((x, y), R, dxfattribs={"layer": LAY, "color": 3})
        t = msp.add_text(str(sn), height=300, dxfattribs={"layer": LAY, "color": 1})
        t.dxf.insert = (x + R * 0.7, y + R * 0.7)
    os.makedirs(os.path.dirname(out_dxf), exist_ok=True)
    doc.saveas(out_dxf)
    return out_dxf

# =================== استخراج آلي من لوحات الأسقف ===================
PLAN_RE = re.compile(r'ADDITIONAL TOP REINFORCEMENT PLAN FOR (.+)', re.I)

def plan_tiles(msp):
    ts = []
    for e in msp.query("TEXT MTEXT"):
        m = PLAN_RE.search(ct.text_of(e))
        if m:
            try: y = float(e.dxf.insert.y)
            except Exception: continue
            ts.append((y, m.group(1).strip().upper()))
    return ts

def _plan_of(y, ts):
    return min(ts, key=lambda p: abs(p[0] - y))[1] if ts else "?"

def detect_levels(msp, ts):
    """منسوب S.S.L لكل سقف من بلوك المنسوب (قد تتعدّد لكل سقف لأن البلاطة مُدرَّجة)."""
    lv = collections.defaultdict(list)
    for ins in msp.query("INSERT"):
        nm = ins.dxf.name.upper()
        if "LVL" not in nm and "LEVEL" not in nm: continue
        try: atts = list(ins.attribs)
        except Exception: atts = []
        for a in atts:
            mm = re.search(r'([-+]?\d{1,2}\.\d{1,2})', " ".join((a.dxf.text or "").split()))
            if mm:
                try: y = float(ins.dxf.insert.y)
                except Exception: continue
                lv[_plan_of(y, ts)].append(float(mm.group(1)))
    return {p: sorted(set(v for v in vs if -12 < v < 12)) for p, vs in lv.items()}

def detect_slab(msp, ts):
    sl = collections.defaultdict(collections.Counter)
    for e in msp.query("TEXT MTEXT"):
        m = re.search(r'SLAB\s*=\s*(\d+)', ct.text_of(e), re.I)
        if m:
            try: y = float(e.dxf.insert.y)
            except Exception: continue
            sl[_plan_of(y, ts)][int(m.group(1))] += 1
    return {p: c.most_common(1)[0][0] for p, c in sl.items()}

def count_drops(msp, ts):
    dp = collections.defaultdict(collections.Counter)
    for e in msp.query("TEXT MTEXT"):
        if e.dxf.layer != "S-TEXT": continue
        t = ct.text_of(e)
        if re.fullmatch(r'[2-4]\d0', t):
            try: y = float(e.dxf.insert.y)
            except Exception: continue
            dp[_plan_of(y, ts)][int(t)] += 1
    return {p: dict(c) for p, c in dp.items()}

# =================== Excel: حصر كل عمود ===================
def build_excel(cols_rows, stage_name, height_note, ceilings_line, grand, waste, out_path):
    NAVY, BLUE, LT, GREY, GRN = "1F3864", "2E5496", "D9E1F2", "F2F2F2", "E2EFDA"
    thin = Side(style="thin", color="BFBFBF"); bd = Border(thin, thin, thin, thin)
    wb = Workbook(); ws = wb.active; ws.title = "حصر خرسانة الأعمدة"; ws.sheet_view.rightToLeft = True
    def st(c, *, b=False, sz=11, color="000000", fill=None, ha="center", wrap=False):
        c.font = Font(bold=b, size=sz, color=color); c.alignment = Alignment(ha, "center", wrap_text=wrap)
        if fill: c.fill = PatternFill("solid", fgColor=fill)
        c.border = bd
    ws.merge_cells("A1:G1"); st(ws["A1"], b=True, sz=14, color="FFFFFF", fill=NAVY)
    ws["A1"] = f"حصر خرسانة الأعمدة — {stage_name}"
    ws.merge_cells("A2:G2"); st(ws["A2"], b=True, sz=11, color="1F3864", fill=GRN, ha="center", wrap=True)
    ws["A2"] = ceilings_line; ws.row_dimensions[2].height = 30
    ws.merge_cells("A3:G3"); st(ws["A3"], sz=9, color="FFFFFF", fill=BLUE)
    ws["A3"] = f"الحجم = الطول × العرض × الارتفاع الصافي | {height_note}"
    hdr = ["رقم العمود", "الماركة", "الطول (م)", "العرض (م)", "الارتفاع (م)", "الحجم (م³)", "ملاحظة"]
    for j, h in enumerate(hdr, 1): st(ws.cell(4, j, h), b=True, fill=NAVY, color="FFFFFF")
    r = 5
    for (sn, m, W, H, hC, vol, note) in cols_rows:
        st(ws.cell(r, 1, sn), b=True, fill=GREY); st(ws.cell(r, 2, m))
        st(ws.cell(r, 3, round(W / 1000.0, 2))); st(ws.cell(r, 4, round(H / 1000.0, 2)))
        st(ws.cell(r, 5, round(hC, 2))); st(ws.cell(r, 6, round(vol, 3)))
        st(ws.cell(r, 7, note or "")); r += 1
    st(ws.cell(r, 1, "الإجمالي"), b=True, fill=NAVY, color="FFFFFF")
    for cc in (2, 3, 4, 5): st(ws.cell(r, cc, None), fill=LT)
    st(ws.cell(r, 6, round(grand, 2)), b=True, fill=LT); st(ws.cell(r, 7, None), fill=LT); r += 1
    st(ws.cell(r, 1, f"بعد الهالك +{waste:.0f}٪"), b=True, fill=GREY)
    for cc in (2, 3, 4, 5, 7): st(ws.cell(r, cc, None), fill=GREY)
    st(ws.cell(r, 6, round(grand * (1 + waste / 100.0), 2)), b=True, fill=GREY); r += 2
    ws.merge_cells(f"A{r}:G{r}")
    st(ws.cell(r, 1, "القاعدة: ارتفاع العمود من سقف الدور الأسفل إلى سقف الدور (سقف↔سقف) ناقص سُمك البلاطة (٢٢سم) أو الـ drop panel (٣٢سم). المقطع من جدول الأعمدة لمرحلة هذا الدور. الأرقام تطابق دوائر لوحة المراجعة."), sz=9, fill=GRN, ha="right", wrap=True); ws.row_dimensions[r].height = 42
    ws.column_dimensions["A"].width = 11
    for cc in "BCDEFG": ws.column_dimensions[cc].width = 15
    os.makedirs(os.path.dirname(out_path), exist_ok=True); wb.save(out_path)

# =================== main ===================
def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--dxf", default=DEFAULT_DXF)
    ap.add_argument("--out", default=DEFAULT_OUT)
    ap.add_argument("--stage", type=int, help="رقم المرحلة/الدور (1..5) المطلوب حصره")
    ap.add_argument("--height", type=float, help="ارتفاع المرحلة سقف↔سقف (م) — مؤقتاً رقم واحد")
    ap.add_argument("--slab", type=float, default=0.22, help="سُمك البلاطة (م) يُخصم (افتراضي 0.22)")
    ap.add_argument("--drop", type=float, default=0.32, help="سُمك الـ drop panel (م) حيث يوجد (افتراضي 0.32)")
    ap.add_argument("--drop-serials", default="", help="أرقام الأعمدة تحتها drop panel (تُخصم 0.32)، مثال: 3,7,11")
    ap.add_argument("--mode", choices=["gross", "net"], default="net")
    ap.add_argument("--waste", type=float, default=5.0)
    ap.add_argument("--dwg", action="store_true", help="حوِّل لوحة المراجعة إلى DWG")
    ap.add_argument("--debug", action="store_true", help="اعرض مقاطع المراحل والاستخراج الآلي فقط")
    args = ap.parse_args()

    doc = ezdxf.readfile(args.dxf); msp = doc.modelspace()
    plan, sched = ct.collect(msp)
    counts = ct.count_marks(plan)
    stages_by_mark, stage_names = parse_schedule_stages(sched)
    cols = columns_list(plan)
    marks = sorted([m for m in counts if not m.startswith("CORE")], key=ct.mark_key)

    # مقاطع المراحل (تحقّق)
    print("=== مقطع كل مرحلة لكل ماركة (مم) ===")
    for m in marks:
        cells = [f"{s[0]}×{s[1]}" if s else "----" for s in stages_by_mark.get(m, [])]
        print(f"  {m:5s} (عدد {counts[m]:2d}): " + " | ".join(f"{c:>9s}" for c in cells))
    print("\nالمراحل:")
    for i, n in enumerate(stage_names, 1): print(f"  {i}. {n}")

    # استخراج آلي
    ts = plan_tiles(msp)
    print("\n=== مُستخرَج آلياً من لوحات الأسقف ===")
    lv, sb, dr = detect_levels(msp, ts), detect_slab(msp, ts), count_drops(msp, ts)
    for _, p in sorted(ts):
        print(f"  {p:26s} | S.S.L={lv.get(p, [])} | SLAB={sb.get(p)}mm | drops={dr.get(p, {})}")

    if args.debug or not args.stage:
        print("\n(لحصر دور: --stage N --height H [--dwg]. المناسيب قد تتعدّد لكل دور — سيُؤخذ ارتفاع لكل عمود عند توفّر المناسيب.)")
        return

    si = args.stage - 1
    if not (0 <= si < len(stage_names)):
        print("رقم مرحلة غير صحيح."); return
    drop_ser = set(int(s) for s in re.findall(r'\d+', args.drop_serials))

    # اسما السقفين ومناسيبهما → ارتفاع تلقائي من المناسيب (الغالب = الأعلى)
    mm = re.match(r'FROM\s+(.+?)\s+TO\s+(.+)', stage_names[si], re.I)
    from_c = mm.group(1).strip() if mm else "?"; to_c = mm.group(2).strip() if mm else "?"
    def levels_of(cname):
        plan = SCHED_TO_PLAN.get(cname.upper(), cname.upper())
        for p, vals in lv.items():
            if plan in p or p in plan: return vals
        return []
    fl, tl = levels_of(from_c), levels_of(to_c)
    dom = lambda vals: max(vals) if vals else None       # المنسوب الغالب (الأعلى)؛ السفلية درجات محلّية
    auto_h = round(dom(tl) - dom(fl), 2) if (fl and tl) else None
    height = args.height if args.height is not None else auto_h
    if height is None:
        print("لا مناسيب كافية لهذا الدور باللوحة — مرّر --height (مناسيب المتكرّر/الروف لم تصل بعد)."); return
    multi = len(fl) > 1 or len(tl) > 1
    fs = f" (منسوب {'/'.join(f'{v:+.2f}' for v in fl)})" if fl else ""
    tsx = f" (منسوب {'/'.join(f'{v:+.2f}' for v in tl)})" if tl else ""
    ceilings_line = f"الحصر بين سقفين — من: {gloss(from_c)}{fs}  ◄──  إلى: {gloss(to_c)}{tsx}"

    # أعمدة هذا الدور فقط، مُرقّمة 1..N (تطابق دوائر اللوحة الإكسل)
    present = []
    for (sn, m, x, y) in cols:
        if m.startswith("CORE"): continue           # الكورات منفصلة
        wh = stages_by_mark.get(m, [None] * len(stage_names))[si] if m in stages_by_mark else None
        if not wh: continue                          # الماركة غير موجودة بهذه المرحلة
        present.append((m, x, y, wh))
    cols_rows = []; review_cols = []; grand = 0.0
    for i, (m, x, y, wh) in enumerate(present, 1):
        ded = (args.drop if i in drop_ser else args.slab) if args.mode == "net" else 0.0
        hC = max(0.0, height - ded)
        vol = section_area_m2(wh) * hC
        note = "drop panel ٣٢سم" if i in drop_ser else ""
        cols_rows.append((i, m, wh[0], wh[1], hC, vol, note))
        review_cols.append((i, m, x, y)); grand += vol

    stage_name = f"المرحلة {args.stage}: {stage_names[si]}"
    hsrc = "تلقائي من المناسيب" if args.height is None else "مُدخَل"
    hnote = (f"ارتفاع الغالب {height}م ({hsrc}، سقف↔سقف − بلاطة/هابط) | {len(cols_rows)} عمود"
             + ("  ⚠ السقف بمنسوب متعدّد (بلاطة مُدرَّجة) → أعمدة المنطقة المنخفضة أعلى ارتفاعاً، تُراجَع" if multi else ""))
    print(f"\n=== {stage_name} ===  ({len(cols_rows)} عمود)")
    print("  " + ceilings_line)
    for (sn, m, W, H, hC, vol, note) in cols_rows[:999]:
        print(f"  #{sn:<3d} {m:5s}  {W/1000:.2f}×{H/1000:.2f}×{hC:.2f} = {vol:.3f} م³ {note}")
    print(f"  ── الإجمالي: {grand:.2f} م³  | بعد الهالك +{args.waste:.0f}٪: {grand*(1+args.waste/100):.2f} م³")

    base = args.out
    xlsx = base + f" - {stage_names[si][:30]}.xlsx"
    build_excel(cols_rows, stage_name, hnote, ceilings_line, grand, args.waste, xlsx)
    print("SAVED:", xlsx)
    review = base + " - review.dxf"
    build_review(args.dxf, review_cols, review)
    print("SAVED:", review)
    if args.dwg and takeoff and hasattr(takeoff, "to_dwg"):
        try: print("DWG:", takeoff.to_dwg(review) or "فشل")
        except Exception as e: print("تعذّر DWG:", e)

if __name__ == "__main__":
    main()
