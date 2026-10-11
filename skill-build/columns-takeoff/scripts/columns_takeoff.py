# -*- coding: utf-8 -*-
"""
حصر أعداد الأعمدة من لوحة COLUMNS & AXES في ملف الـ DXF.
Column COUNT takeoff from a structural COLUMNS & AXES plan.

يقرأ علامات الأعمدة (C, C1..Cn) والكورات (CORE 1..n) من اللوحة، يعدّها حسب الماركة،
يربط كل ماركة بمقطعها وتسليحها من جدول الأعمدة (COLUMN SCHEDULE)، ويخرّج جدول Excel منظّم.

الاستخدام:
    python columns_takeoff.py --dxf "<file.dxf>"
    python columns_takeoff.py --dxf "<file.dxf>" --out "<out.xlsx>"
    python columns_takeoff.py --dxf "<file.dxf>" --debug-render   # صورة مراجعة (نص أسود)

يعتمد على العناوين لتحديد المناطق تلقائياً:
  * لوحة الأعمدة  : أقرب نص فيه "COLUMN" و"AX" (مثل "COLUMNS&AXES").
  * جدول الأعمدة  : نص "COLUMN SCHEDULE".
لو تعذّر إيجادها استخدم --plan-bbox / --sched-bbox يدوياً.
"""
import argparse, os, re, math, collections
import ezdxf
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

# علامة الماركة: C وحدها أو C+أرقام. الكور منفصل.
MARK_RE = re.compile(r"^C\s*\.?\s*(\d{0,3})$", re.I)
CORE_RE = re.compile(r"^CORE\s*\.?\s*(\d+)$", re.I)
SIZE_RE = re.compile(r"SIZE\s*=\s*(\d+)\s*[xX]\s*(\d+)", re.I)
VERT_RE = re.compile(r"VERTICAL\s*=\s*(.+)$", re.I)

def text_of(e):
    try:
        t = e.plain_text() if e.dxftype() == "MTEXT" else e.dxf.text
    except Exception:
        t = getattr(e.dxf, "text", "")
    return " ".join((t or "").split())

def mark_key(m):
    if m == "C":
        return (0, 0)
    mm = re.match(r"C(\d+)$", m)
    if mm:
        return (0, int(mm.group(1)))
    cm = re.match(r"CORE(\d+)$", m)
    if cm:
        return (1, int(cm.group(1)))
    return (2, 0)

def all_texts(msp):
    out = []
    for e in msp.query("TEXT MTEXT"):
        try:
            ip = e.dxf.insert; x, y = float(ip.x), float(ip.y)
        except Exception:
            continue
        t = text_of(e)
        if t:
            out.append((x, y, t))
    return out

def find_title(texts, *needles):
    ns = [n.upper() for n in needles]
    hits = [(x, y, t) for (x, y, t) in texts if all(n in t.upper() for n in ns)]
    # أقصر نص مطابق (العنوان نفسه لا فقرة تحتويه)
    return min(hits, key=lambda h: len(h[2])) if hits else None

def in_box(x, y, box):
    return box[0] <= x <= box[1] and box[2] <= y <= box[3]

def collect(texts, plan_box, sched_box):
    plan = [(x, y, t) for (x, y, t) in texts if in_box(x, y, plan_box)]
    sched = [(x, y, t) for (x, y, t) in texts if in_box(x, y, sched_box)]
    return plan, sched

def count_marks(plan):
    counts = collections.Counter()
    for (x, y, t) in plan:
        mm = MARK_RE.match(t); cm = CORE_RE.match(t)
        if mm:
            counts["C" + mm.group(1)] += 1
        elif cm:
            counts["CORE" + cm.group(1)] += 1
    return counts

def parse_schedule(sched):
    """لكل ماركة: مقطع المرحلة الأولى (أقصى يسار) + التسليح الرأسي.
    العنوان أعلى الخلية والمقطع/التسليح أسفله بنحو 1000 وحدة، فنافذة الصف غير متماثلة."""
    if not sched:
        return {}
    heads = [(x, y, "C" + re.match(r"^C\s*(\d{0,3})$", t, re.I).group(1))
             for (x, y, t) in sched if re.match(r"^C\s*\d{0,3}$", t, re.I)]
    if not heads:
        return {}
    head_x = min(h[0] for h in heads)                 # عمود الماركات على أقصى يسار الجدول
    heads = [h for h in heads if abs(h[0] - head_x) < 1500]
    info = {}
    for (hx, hy, mark) in heads:
        row = [(x, y, t) for (x, y, t) in sched if hy - 1450 < y < hy + 250]
        sizes = sorted([(x, SIZE_RE.search(t)) for (x, y, t) in row if SIZE_RE.search(t)])
        verts = sorted([(x, VERT_RE.search(t)) for (x, y, t) in row if VERT_RE.search(t)])
        size = f"{sizes[0][1].group(1)}×{sizes[0][1].group(2)}" if sizes else "—"
        vert = verts[0][1].group(1).strip() if verts else "—"
        info[mark] = (size, vert)
    return info

def build_excel(counts, info, title, path):
    marks = sorted([m for m in counts if not m.startswith("CORE")], key=mark_key)
    cores = sorted([m for m in counts if m.startswith("CORE")], key=mark_key)
    wb = Workbook(); ws = wb.active; ws.title = "حصر الأعمدة"; ws.sheet_view.rightToLeft = True
    NAVY, BLUE, LT, GREY = "1F3864", "2E5496", "D9E1F2", "F2F2F2"
    thin = Side(style="thin", color="BFBFBF"); bd = Border(thin, thin, thin, thin)
    def st(c, *, b=False, sz=11, color="000000", fill=None, ha="center", wrap=False):
        c.font = Font(bold=b, size=sz, color=color, name="Calibri")
        c.alignment = Alignment(horizontal=ha, vertical="center", wrap_text=wrap)
        if fill: c.fill = PatternFill("solid", fgColor=fill)
        c.border = bd
    ws.merge_cells("A1:D1"); st(ws["A1"], b=True, sz=14, color="FFFFFF", fill=NAVY)
    ws["A1"] = f"حصر الأعمدة — {title}"
    ws.merge_cells("A2:D2"); st(ws["A2"], sz=9, color="FFFFFF", fill=BLUE)
    ws["A2"] = "عدد الأعمدة حسب الماركة | المقطع والتسليح من جدول الأعمدة (المرحلة الأولى: الأساس ← سقف البدروم العلوي)"
    ws.row_dimensions[1].height = 26
    for j, h in enumerate(["ماركة العمود", "العدد", "المقطع (مم)", "التسليح الرأسي"], 1):
        st(ws.cell(4, j, h), b=True, fill=NAVY, color="FFFFFF")
    r = 5; total = 0
    for m in marks:
        size, vert = info.get(m, ("—", "—"))
        st(ws.cell(r, 1, m), b=True, fill=GREY); st(ws.cell(r, 2, counts[m]), b=True)
        st(ws.cell(r, 3, size)); st(ws.cell(r, 4, vert)); total += counts[m]; r += 1
    st(ws.cell(r, 1, "إجمالي الأعمدة"), b=True, fill=NAVY, color="FFFFFF")
    st(ws.cell(r, 2, total), b=True, fill=LT); st(ws.cell(r, 3, "")); st(ws.cell(r, 4, "")); r += 2
    if cores:
        ws.merge_cells(f"A{r}:D{r}"); st(ws.cell(r, 1, "الكورات (حوائط قص) — Cores"), b=True, fill=BLUE, color="FFFFFF"); r += 1
        ctot = 0
        for m in cores:
            st(ws.cell(r, 1, m.replace("CORE", "CORE "))); st(ws.cell(r, 2, counts[m]))
            st(ws.cell(r, 3, "")); st(ws.cell(r, 4, "")); ctot += counts[m]; r += 1
        st(ws.cell(r, 1, "إجمالي الكورات"), b=True, fill=NAVY, color="FFFFFF")
        st(ws.cell(r, 2, ctot), b=True, fill=LT); st(ws.cell(r, 3, "")); st(ws.cell(r, 4, "")); r += 2
    ws.merge_cells(f"A{r}:D{r}")
    st(ws.cell(r, 1, "ملاحظة: «C» ماركة مستقلة عن «C1» (راجع جدول الأعمدة). المقاطع للمرحلة الأولى وتصغر بالأدوار الأعلى."),
       sz=9, color="1F3864", fill="E2EFDA", ha="right", wrap=True); ws.row_dimensions[r].height = 34
    for col, w in zip("ABCD", (18, 10, 16, 18)):
        ws.column_dimensions[col].width = w
    os.makedirs(os.path.dirname(os.path.abspath(path)), exist_ok=True)
    wb.save(path)

def debug_render(dxf, plan_box, out_png):
    """صورة مراجعة: النصوص مرسومة سوداء (علامات الأعمدة لونها ACI 7 أبيض فلا تظهر افتراضياً)."""
    import matplotlib; matplotlib.use("Agg")
    from ezdxf.addons.drawing import RenderContext, Frontend
    from ezdxf.addons.drawing.matplotlib import MatplotlibBackend
    import matplotlib.pyplot as plt
    doc = ezdxf.readfile(dxf); msp = doc.modelspace()
    ents = []
    for e in msp:
        tp = e.dxftype()
        try:
            if tp in ("TEXT", "MTEXT"):
                p = e.dxf.insert; x, y = float(p.x), float(p.y)
            elif tp == "CIRCLE":
                x, y = float(e.dxf.center.x), float(e.dxf.center.y)
            elif tp == "LINE":
                x = (e.dxf.start.x + e.dxf.end.x) / 2; y = (e.dxf.start.y + e.dxf.end.y) / 2
            elif tp in ("LWPOLYLINE", "POLYLINE"):
                pts = list(e.get_points("xy")) if tp == "LWPOLYLINE" else None
                if not pts: continue
                x = sum(p[0] for p in pts) / len(pts); y = sum(p[1] for p in pts) / len(pts)
            else:
                continue
        except Exception:
            continue
        if in_box(x, y, plan_box):
            if tp in ("TEXT", "MTEXT"):
                try: e.rgb = (0, 0, 0)
                except Exception: pass
            ents.append(e)
    fig = plt.figure(figsize=(20, 16)); ax = fig.add_axes([0, 0, 1, 1])
    Frontend(RenderContext(doc), MatplotlibBackend(ax)).draw_entities(ents)
    ax.set_xlim(plan_box[0], plan_box[1]); ax.set_ylim(plan_box[2], plan_box[3])
    ax.set_aspect("equal"); ax.axis("off")
    fig.savefig(out_png, dpi=180, facecolor="white"); plt.close(fig)

def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--dxf", required=True)
    ap.add_argument("--out", default=None)
    ap.add_argument("--plan-bbox", nargs=4, type=float, metavar=("XMIN", "XMAX", "YMIN", "YMAX"))
    ap.add_argument("--sched-bbox", nargs=4, type=float, metavar=("XMIN", "XMAX", "YMIN", "YMAX"))
    ap.add_argument("--debug-render", action="store_true")
    args = ap.parse_args()

    doc = ezdxf.readfile(args.dxf); msp = doc.modelspace()
    texts = all_texts(msp)

    # منطقة اللوحة من عنوان COLUMNS&AXES، والجدول من COLUMN SCHEDULE
    if args.plan_bbox:
        plan_box = tuple(args.plan_bbox)
    else:
        t = find_title(texts, "COLUMN", "AX")           # COLUMNS&AXES
        if not t:
            raise SystemExit("تعذّر إيجاد عنوان لوحة الأعمدة؛ مرّر --plan-bbox يدوياً.")
        tx, ty, _ = t
        plan_box = (tx - 20000, tx + 45000, ty + 6000, ty + 62000)
    if args.sched_bbox:
        sched_box = tuple(args.sched_bbox)
    else:
        s = find_title(texts, "COLUMN SCHEDULE")
        if s:
            sx, sy, _ = s
            sched_box = (sx - 3000, sx + 45000, sy - 25000, sy + 3000)
        else:
            sched_box = (1e9, 1e9, 1e9, 1e9)            # لا جدول

    plan, sched = collect(texts, plan_box, sched_box)
    counts = count_marks(plan)
    info = parse_schedule(sched)

    marks = sorted([m for m in counts if not m.startswith("CORE")], key=mark_key)
    cores = sorted([m for m in counts if m.startswith("CORE")], key=mark_key)
    print(f"plan_box={tuple(round(v) for v in plan_box)}  marks={len(plan)} texts")
    print("=== عدد الأعمدة ===")
    tot = 0
    for m in marks:
        size, vert = info.get(m, ("—", "—"))
        print(f"  {m:5s} : {counts[m]:3d}   [{size}  {vert}]"); tot += counts[m]
    print(f"  إجمالي الأعمدة: {tot}")
    if cores:
        print("=== الكورات ===")
        for m in cores:
            print(f"  {m.replace('CORE','CORE ')} : {counts[m]}")

    out = args.out or os.path.join(os.path.dirname(os.path.abspath(args.dxf)),
                                   "output", "حصر الأعمدة - COLUMNS & AXES.xlsx")
    build_excel(counts, info, "COLUMNS & AXES", out)
    print("\nSAVED:", out)
    if args.debug_render:
        png = os.path.splitext(out)[0] + " - review.png"
        debug_render(args.dxf, plan_box, png)
        print("SAVED:", png)

if __name__ == "__main__":
    main()
