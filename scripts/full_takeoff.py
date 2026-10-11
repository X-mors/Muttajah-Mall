# -*- coding: utf-8 -*-
"""
حصر حديد كامل لسقف واحد = التسليح الإضافي (علوي+سفلي) + الفرش والغطاء (الشبكة الرئيسية)
في شيت إكسل واحد منظّم، مع لوحة إثبات (دوائر مرقّمة للإضافي + الفتحات المكتشفة).

python full_takeoff.py "LOWER BASEMENT CEILING" [--dwg] [--openings <m²>] [--gx <w> --gy <l>]
"""
import argparse, os, math, collections, json
import ezdxf
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import takeoff, slab_mesh

def detect_openings_full(msp, box):
    """يفوّض لكاشف الفتحات في slab_mesh (تقاطع قُطرَي كل X → فصل الفتحات المتلاصقة)؛ يعيد قائمة (م²,cx,cy,w,h)."""
    _, items = slab_mesh.detect_openings(msp, box)
    return items

def mesh_notes(msp, tbox, bbox):
    notes = slab_mesh.region_notes(msp, tbox) + (slab_mesh.region_notes(msp, bbox) if bbox else [])
    seen = set(); out = []
    for n in notes:
        k = (n["face"], n["dia"], n["sp"])
        if k in seen: continue
        seen.add(k); out.append(n)
    return out

def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("ceiling"); ap.add_argument("--dxf", default=takeoff.DEFAULT_DXF)
    ap.add_argument("--outdir", default=takeoff.DEFAULT_OUT); ap.add_argument("--dwg", action="store_true")
    ap.add_argument("--openings", type=float); ap.add_argument("--gx", type=float); ap.add_argument("--gy", type=float)
    ap.add_argument("--cut", action="store_true", help="شغّل تفريد الحديد (rebar-cutting) وأصدِر جدول التفريد")
    ap.add_argument("--lap", type=int, default=60, help="مضاعف طول وصلة التراكب (×القطر): للشبكة الرئيسية، وللتفريد >12م. 60 شد/40 ضغط")
    ap.add_argument("--conservative", action="store_true", help="احتسب قصاصة آخر سيخ لكل مجرى شبكة. الافتراضي: تراكب فقط (محسّن)")
    ap.add_argument("--include-mesh", dest="include_mesh", action="store_true", help="أظهِر الشبكة في جدول التفريد")
    args = ap.parse_args()
    args.reuse = not args.conservative   # الافتراضي = محسّن (تراكب فقط لمجاري الشبكة)
    doc = ezdxf.readfile(args.dxf); msp = doc.modelspace(); titles = takeoff.all_titles(msp)
    th = takeoff.find_title(titles, args.ceiling, "TOP")
    if not th: print("لا توجد لوحة TOP لهذا السقف."); return
    tx, ty, ttl = th[0]; tbox = takeoff.plan_bbox(titles, tx, ty)
    bh = takeoff.find_title(titles, args.ceiling, "BOTTOM"); bbox = None
    if bh: bbox = takeoff.plan_bbox(titles, bh[0][0], bh[0][1])

    # ---------- ADDITIONAL (top + bottom) ----------
    allc = []; serial = 0
    for label, box in [("علوي", tbox)] + ([("سفلي", bbox)] if bbox else []):
        cs = takeoff.extract(msp, box)
        for c in sorted(cs, key=lambda c: (-c["y"], c["x"])):
            serial += 1; c["serial"] = serial; c["_plan"] = label; allc.append(c)
    add_piv = collections.defaultdict(lambda: collections.defaultdict(int))  # dia -> L -> count
    add_len = collections.defaultdict(float); add_tot = collections.defaultdict(int)
    n_stair = 0
    for c in allc:
        if c["stair"]: n_stair += 1; continue
        if not c["count"]: continue
        L = c["L"] or 0
        add_piv[c["dia"]][L] += c["count"]; add_tot[c["dia"]] += c["count"]
        add_len[c["dia"]] += c["count"] * (L / 1000.0)

    # تفريد الإضافي: أسياخه الفعلية وفائضه (قطع كاملة قابلة للاستخدام في الشبكة لنفس القطر)
    try:
        import cutlist
        _cut = cutlist.optimize(add_piv)
    except Exception:
        _cut = None
    add_stock = ({d: _cut[d]["stock_bars"] for d in _cut} if _cut
                 else {d: math.ceil(add_len[d] / 12.0) for d in add_len})
    add_surplus_m = {d: (_cut[d]["surplus_mm"] / 1000.0 if _cut else 0.0) for d in add_piv}
    add_surplus_n = {d: (sum(_cut[d]["surplus"].values()) if _cut else 0) for d in add_piv}

    # ---------- MAIN MESH ----------
    gx = args.gx or slab_mesh.rebar_extent(msp, tbox)[0]
    gy = args.gy or slab_mesh.rebar_extent(msp, tbox)[1]
    gross = gx * gy
    ops = detect_openings_full(msp, tbox)
    openings = args.openings if args.openings is not None else sum(o[0] for o in ops)
    net = gross - openings
    notes = mesh_notes(msp, tbox, bbox)
    mesh_len = collections.defaultdict(float)
    for nn in notes: mesh_len[nn["dia"]] += net * (2000.0 / nn["sp"])
    # زيادة الوصلات للشبكة = وصلات تراكب حقيقية (lap×القطر) بدل +10٪ المؤقتة السابقة
    mesh_f = {d: slab_mesh.lap_factor(gx, gy, d, args.lap, args.reuse) for d in mesh_len}
    mesh_lenw = {d: mesh_len[d] * mesh_f[d] for d in mesh_len}
    mesh_bars = {d: math.ceil(mesh_lenw[d] / 12.0) for d in mesh_len}   # تصميمي (مستقل)

    # ---------- grand total meters per diameter ----------
    dias = sorted(set(list(add_len) + list(mesh_len)))
    grand_m = {d: round(add_len.get(d, 0) + mesh_lenw.get(d, 0)) for d in dias}  # تصميمي (لا يتغيّر)
    # فائض الإضافي (قطع كاملة) يُستعمل في الشبكة لنفس القطر: كل قطعة توفّر (طولها − وصلة) من صلب الشبكة
    mesh_reuse_bars = {}; mesh_buy = {}
    for d in mesh_len:
        lap_m = max(args.lap * d / 1000.0, 1.0)
        reuse_m = max(0.0, add_surplus_m.get(d, 0.0) - add_surplus_n.get(d, 0) * lap_m)
        reuse_m = min(reuse_m, mesh_lenw[d])
        mesh_buy[d] = math.ceil((mesh_lenw[d] - reuse_m) / 12.0)
        mesh_reuse_bars[d] = mesh_bars[d] - mesh_buy[d]
    # وزن الشراء = أسياخ الإضافي الفعلية + أسياخ الشبكة بعد خصم الفائض المُستعمَل
    buy_bars = {d: add_stock.get(d, 0) + mesh_buy.get(d, 0) for d in dias}
    buy_w = {d: round(buy_bars[d] * 12.0 * d * d / 162.0) for d in dias}

    print(f"السقف: {ttl}")
    print("الإضافي (م/قطر):", {f"Ø{d}": round(add_len[d]) for d in add_len}, "| بنود:", sum(add_tot.values()))
    print(f"الشبكة: صافي {net:.0f}م² | مواصفات:", [f"{n['face']} T{n['dia']}@{n['sp']}" for n in notes])
    print(f"الشبكة (م/قطر، وصلات {args.lap}×قطر{' تراكب-فقط' if args.reuse else ''}):",
          {f"Ø{d}": f"{round(mesh_lenw[d])} (×{mesh_f[d]:.3f})" for d in mesh_lenw})
    print("الإجمالي الكلي (م/قطر):", {f"Ø{d}": grand_m[d] for d in dias})

    # ---------- EXCEL ----------
    NAVY,BLUE,LT,GREY,GRN = "1F3864","2E5496","D9E1F2","F2F2F2","E2EFDA"
    thin = Side(style="thin", color="BFBFBF"); bd = Border(thin, thin, thin, thin)
    wb = Workbook()
    def st(c,*,b=False,sz=11,color="000000",fill=None,ha="center",wrap=False):
        c.font=Font(bold=b,size=sz,color=color); c.alignment=Alignment(ha,"center",wrap_text=wrap)
        if fill: c.fill=PatternFill("solid",fgColor=fill)
        c.border=bd
    ws = wb.active; ws.title="الملخص"; ws.sheet_view.rightToLeft=True
    ws.merge_cells("A1:F1"); st(ws["A1"],b=True,sz=15,color="FFFFFF",fill=NAVY); ws["A1"]=f"الحصر الكامل — {ttl}"
    r=3
    # (A) additional
    ws.merge_cells(f"A{r}:F{r}"); st(ws.cell(r,1,"◄ أولاً: التسليح الإضافي (علوي+سفلي) — العدد لكل مقاس + الوزن"),b=True,color="FFFFFF",fill=BLUE,ha="right"); r+=1
    adias=sorted(add_piv); st(ws.cell(r,1,"طول القطعية (م)"),b=True,fill=NAVY,color="FFFFFF")
    for j,d in enumerate(adias): st(ws.cell(r,2+j,f"Ø{d}"),b=True,fill=NAVY,color="FFFFFF")
    r+=1
    Ls=sorted({L for d in add_piv for L in add_piv[d]})
    for L in Ls:
        st(ws.cell(r,1, f"{L/1000:g}" if L else "—"),fill=GREY)
        for j,d in enumerate(adias): st(ws.cell(r,2+j, add_piv[d].get(L) or None))
        r+=1
    st(ws.cell(r,1,"الوزن (كجم)"),b=True,fill=LT)
    for j,d in enumerate(adias): st(ws.cell(r,2+j, round(add_len[d]*d*d/162.0) or None),b=True,fill=LT)
    r+=1
    addw=round(sum(add_len[d]*d*d/162.0 for d in adias))
    st(ws.cell(r,1,"إجمالي وزن الإضافي (كجم)"),b=True,fill=LT)
    ws.merge_cells(f"B{r}:{get_column_letter(1+len(adias))}{r}"); st(ws.cell(r,2, addw),b=True,fill=LT)
    r+=2
    # (B) mesh
    ws.merge_cells(f"A{r}:F{r}"); st(ws.cell(r,1,"◄ ثانياً: الفرش والغطاء (الشبكة الرئيسية) — منهج المساحة"),b=True,color="FFFFFF",fill=BLUE,ha="right"); r+=1
    ws.merge_cells(f"A{r}:F{r}"); st(ws.cell(r,1,f"إجمالي {gx:.2f}×{gy:.2f}={gross:.0f}م² − فتحات(شاملة السلالم) {openings:.0f} = صافي {net:.0f}م² | وصلات تراكب {args.lap}×القطر {'(تراكب فقط)' if args.reuse else '(قياسي)'} ÷12"),sz=9,fill=GRN,ha="right"); r+=1
    for h,c in [("القطر",1),("صافي الطول (م)",2),("مع الوصلات (م)",3),("عدد أسياخ 12م",4),("المواصفة",5)]:
        st(ws.cell(r,c,h),b=True,fill=NAVY,color="FFFFFF")
    r+=1
    spec_by_d=collections.defaultdict(list)
    for n in notes: spec_by_d[n["dia"]].append(f"{n['face']} T{n['dia']}@{n['sp']}")
    for d in sorted(mesh_len):
        st(ws.cell(r,1,f"Ø{d}"),b=True,fill=LT); st(ws.cell(r,2,round(mesh_len[d]))); st(ws.cell(r,3,round(mesh_lenw[d])))
        st(ws.cell(r,4,mesh_bars[d]),b=True); st(ws.cell(r,5," ، ".join(sorted(set(spec_by_d[d]))))); r+=1
    r+=1
    # (C) grand total meters
    ws.merge_cells(f"A{r}:F{r}"); st(ws.cell(r,1,"◄ الإجمالي الكلي لكل قطر (متر طولي = إضافي + شبكة)"),b=True,color="FFFFFF",fill=NAVY,ha="right"); r+=1
    st(ws.cell(r,1,"القطر"),b=True,fill=BLUE,color="FFFFFF")
    for j,d in enumerate(dias): st(ws.cell(r,2+j,f"Ø{d}"),b=True,fill=BLUE,color="FFFFFF")
    r+=1
    st(ws.cell(r,1,"متر طولي",),b=True,fill=GREY)
    for j,d in enumerate(dias): st(ws.cell(r,2+j,grand_m[d]),b=True)
    r+=1
    st(ws.cell(r,1,"الوزن (كجم)"),b=True,fill=GREY)
    for j,d in enumerate(dias): st(ws.cell(r,2+j, round(grand_m[d]*d*d/162.0)),b=True)
    r+=1
    totw=round(sum(grand_m[d]*d*d/162.0 for d in dias))
    st(ws.cell(r,1,"وزن تصميمي (كجم)",),b=True,fill=GREY)
    ws.merge_cells(f"B{r}:{get_column_letter(1+len(dias))}{r}")
    st(ws.cell(r,2, f"{totw}  ({round(totw/1000.0,2)} طن)"),b=True,fill=LT)
    r+=1
    # فائض الإضافي المُستعمَل في الشبكة لنفس القطر (توفير في الشراء)
    if any(mesh_reuse_bars.get(d,0) for d in dias):
        st(ws.cell(r,1,"فائض إضافي يُستعمَل بالشبكة (سيخ)"),b=True,fill=GRN)
        for j,d in enumerate(dias): st(ws.cell(r,2+j, (f"−{mesh_reuse_bars[d]}" if mesh_reuse_bars.get(d,0) else None)))
        r+=1
    # وزن الشراء (شامل هدر القطع، بعد استعمال الفائض بالشبكة)
    st(ws.cell(r,1,"أسياخ 12م (شراء نهائي)"),b=True,fill=BLUE,color="FFFFFF")
    for j,d in enumerate(dias): st(ws.cell(r,2+j, buy_bars[d]),b=True)
    r+=1
    st(ws.cell(r,1,"وزن الشراء (كجم)"),b=True,fill=BLUE,color="FFFFFF")
    for j,d in enumerate(dias): st(ws.cell(r,2+j, buy_w[d]),b=True)
    r+=1
    buyw=sum(buy_w.values())
    st(ws.cell(r,1,"وزن الشراء الكلي",),b=True,fill=NAVY,color="FFFFFF")
    ws.merge_cells(f"B{r}:{get_column_letter(1+len(dias))}{r}")
    st(ws.cell(r,2, f"{buyw} كجم  ({round(buyw/1000.0,2)} طن) — شامل الهدر، بعد استعمال الفائض بالشبكة"),b=True,fill=LT)
    r+=2
    reuse_note = ("فائض الإضافي (قطع كاملة، نفس القطر) يُستعمَل في الشبكة فيقلّل شراءها (كل قطعة تُوصَّل بوصلة). "
                  if any(mesh_reuse_bars.get(d,0) for d in dias) else "")
    ws.merge_cells(f"A{r}:F{r}"); st(ws.cell(r,1,f"«تصميمي» = أطوال صافية + وصلات الشبكة. «الشراء النهائي» = أسياخ الإضافي (بالتفريد) + أسياخ الشبكة بعد خصم الفائض المُستعمَل فيها. {reuse_note}راجع شيت تفاصيل الفائض والهدر. سلالم مستبعدة = {n_stair} كولّ. الفتحات مخصومة."),sz=9,fill=GRN,ha="right",wrap=True); ws.row_dimensions[r].height=48
    ws.column_dimensions["A"].width=18
    for cc in range(2,8): ws.column_dimensions[get_column_letter(cc)].width=13

    # detail sheets
    ws2=wb.create_sheet("تفاصيل الإضافي"); ws2.sheet_view.rightToLeft=True
    for j,h in enumerate(["الرقم","اللوحة","القطر","التباعد","الوجه","الطول L","العرض","العدد","النوع","x","y"],1): st(ws2.cell(1,j,h),b=True,fill=NAVY,color="FFFFFF")
    for i,c in enumerate(sorted(allc,key=lambda c:c["serial"]),1):
        typ="سلّم" if c["stair"] else "بلاطة"; fill=("E7E6E6" if c["stair"] else (None if c.get("src") in ("circle","barline") else "FFF2CC"))
        row=[c["serial"],c["_plan"],f"Ø{c['dia']}",c["sp"],("علوي" if c["face"]=="T" else "سفلي"),c["L"],c["w"],(None if c["stair"] else c["count"]),typ,round(c["x"]),round(c["y"])]
        for j,v in enumerate(row,1): st(ws2.cell(i+1,j,v),fill=fill)
    for j,wd in enumerate([6,7,7,8,7,9,8,7,9,10,10],1): ws2.column_dimensions[get_column_letter(j)].width=wd
    ws2.freeze_panes="A2"
    ws3=wb.create_sheet("الفتحات"); ws3.sheet_view.rightToLeft=True
    for j,h in enumerate(["الرقم","المساحة (م²)","العرض(مم)","الطول(مم)","x","y"],1): st(ws3.cell(1,j,h),b=True,fill=NAVY,color="FFFFFF")
    for i,(a2,cx,cy,w,h) in enumerate(ops,1):
        for j,v in enumerate([i,round(a2,1),round(w),round(h),round(cx),round(cy)],1): st(ws3.cell(i+1,j,v))
    st(ws3.cell(len(ops)+2,1,"الإجمالي"),b=True,fill=LT); st(ws3.cell(len(ops)+2,2,round(sum(o[0] for o in ops),1)),b=True,fill=LT)
    for j,wd in enumerate([6,12,10,10,10,10],1): ws3.column_dimensions[get_column_letter(j)].width=wd

    os.makedirs(args.outdir,exist_ok=True)
    base=os.path.join(args.outdir, takeoff.safe(ttl)+" - FULL")
    xlsx=base+".xlsx"; wb.save(xlsx)

    # ---------- PIECE-LIST SIDECAR (for the rebar-cutting skill) ----------
    spec_by_dia=collections.defaultdict(list)
    for n in notes: spec_by_dia[n["dia"]].append(f"{n['face']} T{n['dia']}@{n['sp']}")
    piece_list=[{"diameter":d,"length_mm":L,"count":add_piv[d][L]}
                for d in sorted(add_piv) for L in sorted(add_piv[d]) if L and add_piv[d][L]]
    unknown=[{"diameter":d,"count":add_piv[d].get(0,0)} for d in sorted(add_piv) if add_piv[d].get(0)]
    mesh_json=[{"diameter":d,"full_bars_12m":mesh_bars[d],
                "spec":" ، ".join(sorted(set(spec_by_dia[d])))} for d in sorted(mesh_bars)]
    sidecar={"ceiling":ttl,"stock_length_mm":12000,"unit":"mm",
             "pieces":piece_list,"unknown_length":unknown,"mesh":mesh_json}
    pj=base+" - pieces.json"
    with open(pj,"w",encoding="utf-8") as f: json.dump(sidecar,f,ensure_ascii=False,indent=1)
    print("SAVED:",pj)
    if args.cut:
        try:
            import cutlist
            cutmeta={"ceiling":ttl,"stock_length_mm":12000,"unknown_length":unknown,"mesh":mesh_json}
            res,cutout=cutlist.cut_from_takeoff(add_piv,cutmeta,base+" - CUTLIST.xlsx",
                                                lap_mult=args.lap,include_mesh=args.include_mesh)
            cutlist._print_summary(res); print("CUTLIST:",cutout)
        except Exception as e:
            print("تعذّر تشغيل التفريد (--cut):",e,"| استخدم مهارة rebar-cutting على",pj)

    # ---------- REVIEW DRAWING ----------
    doc2=ezdxf.readfile(args.dxf); msp2=doc2.modelspace()
    LAY="CHECK-FULL"
    if LAY not in doc2.layers: doc2.layers.add(LAY,color=7)
    R=850; GREEN,CYAN,RED,MAG,ORANGE=3,4,1,6,30
    for c in allc:
        col = MAG if c["stair"] else (RED if c["count"] is None else (CYAN if c.get("src")=="heur" else GREEN))
        msp2.add_circle((c["x"],c["y"]),R,dxfattribs={"layer":LAY,"color":col})
        t=msp2.add_text(str(c["serial"]),height=300,dxfattribs={"layer":LAY,"color":col}); t.dxf.insert=(c["x"]+R*0.7,c["y"]+R*0.7)
    for i,(a2,cx,cy,w,h) in enumerate(ops,1):
        msp2.add_lwpolyline([(cx-w/2,cy-h/2),(cx+w/2,cy-h/2),(cx+w/2,cy+h/2),(cx-w/2,cy+h/2),(cx-w/2,cy-h/2)],dxfattribs={"layer":LAY,"color":ORANGE})
        t=msp2.add_text(f"OP{i}",height=350,dxfattribs={"layer":LAY,"color":ORANGE}); t.dxf.insert=(cx,cy)
    review=base+" - review.dxf"; doc2.saveas(review)
    print("SAVED:",xlsx); print("SAVED:",review)
    if args.dwg:
        d=takeoff.to_dwg(review); print("DWG:",d or "فشل")

if __name__=="__main__":
    main()
