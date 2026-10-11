# -*- coding: utf-8 -*-
"""
حصر الفرش والغطاء (main slab mesh) بمنهج المساحة/التوريد:
  صافي المساحة = الإجمالي(الطول×العرض) − الفتحات − السلالم
  طول الصلب/م² لشبكة ثنائية بتباعد s(مم) لكل طبقة = 2000/s م/م²
  إجمالي الطول(لكل قطر) = صافي × Σ(2000/s)  ؛  يُضرب في معامل وصلات التراكب (lap×القطر حسب الأبعاد) ثم ÷12

الاستخدام:
  python slab_mesh.py "LOWER BASEMENT CEILING"                 # أبعاد تلقائية من اللوحة
  python slab_mesh.py "LOWER BASEMENT CEILING" --gx 53.9 --gy 45.0   # أبعاد كلّية يدوية (م)
"""
import argparse, os, re, math, collections
import ezdxf
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import takeoff

MESHNOTE = re.compile(r'(TOP|BOT)\.?\s*=\s*T(\d{1,2})@(\d{2,4})', re.I)
STOCK = 12.0          # طول السيخ القياسي (م)
LAP_MIN = 1.0         # أدنى طول وصلة (م) — حدّ عملي حدّده المهندس (١ متر)
DEF_LAP_MULT = 60     # مضاعف طول الوصلة الافتراضي (×القطر، شد). ضغط=40. راجع مهارة rebar-cutting

def in_box(x, y, box): return box[0] <= x <= box[1] and box[2] <= y <= box[3]

def _contrib(run_m, lap_m, reuse):
    """طول الصلب (م) المستهلك لتغطية مجرى شبكة طوله run_m من أسياخ 12م بوصلة lap_m."""
    if run_m <= STOCK: return run_m                      # سيخ واحد، القصاصة تُعاد لمجرى مجاور
    if reuse: return run_m * STOCK / (STOCK - lap_m)     # تراكب فقط (إعادة استخدام القصاصة)
    return math.ceil((run_m - lap_m) / (STOCK - lap_m)) * STOCK   # قياسي: يشمل قصاصة آخر سيخ

def lap_factor(gx, gy, d_mm, lap_mult=DEF_LAP_MULT, reuse=False):
    """معامل زيادة طول الصلب للشبكة بسبب وصلات التراكب — يعتمد على أبعاد البلاطة وقطر السيخ.
    الشبكة ثنائية: مجارٍ بطول gx (وأخرى بطول gy)؛ التباعد يُختصر فلا يظهر. بلا أبعاد → افتراض مجارٍ طويلة."""
    lap_m = max(lap_mult * d_mm / 1000.0, LAP_MIN)
    if not gx or not gy:
        return STOCK / (STOCK - lap_m)                   # لا أبعاد: افترض كل المجاري > 12م (تراكب فقط)
    num = gy * _contrib(gx, lap_m, reuse) + gx * _contrib(gy, lap_m, reuse)
    den = 2.0 * gx * gy
    return num / den if den else 1.0

def region_notes(msp, box):
    out = []
    for ins in msp.query("INSERT"):
        try: atts = list(ins.attribs)
        except Exception: atts = []
        for a in atts:
            m = MESHNOTE.search(" ".join((a.dxf.text or "").split()))
            if not m: continue
            try: x, y = float(a.dxf.insert.x), float(a.dxf.insert.y)
            except Exception: continue
            if in_box(x, y, box):
                out.append(dict(face=m.group(1).upper(), dia=int(m.group(2)), sp=int(m.group(3))))
    return out

def area_sum(msp, box, layer_key):
    tot = 0.0
    for e in msp.query("LWPOLYLINE"):
        if layer_key not in e.dxf.layer.upper(): continue
        try: pts = [(p[0], p[1]) for p in e.get_points("xy")]
        except Exception: continue
        if len(pts) < 3: continue
        cx = sum(p[0] for p in pts) / len(pts); cy = sum(p[1] for p in pts) / len(pts)
        if not in_box(cx, cy, box): continue
        a = 0.0; n = len(pts)
        for i in range(n):
            x1, y1 = pts[i]; x2, y2 = pts[(i + 1) % n]; a += x1 * y2 - x2 * y1
        tot += abs(a) / 2.0 / 1e6
    return tot

def rebar_extent(msp, box):
    """أبعاد البلاطة = المستطيل المحيط بخطوط الحديد ST-RFT-* داخل اللوحة (م)."""
    xs = []; ys = []
    for e in msp.query("LWPOLYLINE LINE"):
        if "RFT" not in e.dxf.layer.upper(): continue
        try:
            ps = [(e.dxf.start.x, e.dxf.start.y), (e.dxf.end.x, e.dxf.end.y)] if e.dxftype() == "LINE" \
                 else [(p[0], p[1]) for p in e.get_points("xy")]
        except Exception: continue
        for x, y in ps:
            if in_box(x, y, box): xs.append(x); ys.append(y)
    if not xs: return 0.0, 0.0
    return (max(xs) - min(xs)) / 1000.0, (max(ys) - min(ys)) / 1000.0

def overall_dims(msp, box, pad=16000):
    """أكبر بُعد أفقي = العرض، أكبر بُعد رأسي = الطول (م). الأبعاد الكلّية تُرسم خارج
    حافّة اللوحة قليلاً، لذا نوسّع النطاق بمقدار pad. نتجاهل الصغيرة والضخمة الشاذّة."""
    big = (box[0] - pad, box[1] + pad, box[2] - pad, box[3] + pad)
    H = V = 0.0
    for e in msp.query("DIMENSION"):
        try:
            p2 = e.dxf.defpoint2; p3 = e.dxf.defpoint3
            ax, ay, bx, by = float(p2.x), float(p2.y), float(p3.x), float(p3.y)
        except Exception: continue
        mx, my = (ax + bx) / 2, (ay + by) / 2
        if not in_box(mx, my, big): continue
        try: v = float(e.get_measurement())
        except Exception: v = 0
        if v <= 5000 or v > 120000: continue   # تجاهل الأبعاد الجزئية والشاذّة
        if abs(bx - ax) >= abs(by - ay): H = max(H, v)
        else: V = max(V, v)
    return H / 1000.0, V / 1000.0

def _seg_cross(a1, b1, a2, b2):
    """نقطة تقاطع قطعتين مستقيمتين إن تقاطعتا داخلياً (0.05≤t,u≤0.95)، وإلا None."""
    d1x, d1y = b1[0] - a1[0], b1[1] - a1[1]
    d2x, d2y = b2[0] - a2[0], b2[1] - a2[1]
    den = d1x * d2y - d1y * d2x
    if abs(den) < 1e-6: return None
    t = ((a2[0] - a1[0]) * d2y - (a2[1] - a1[1]) * d2x) / den
    u = ((a2[0] - a1[0]) * d1y - (a2[1] - a1[1]) * d1x) / den
    if 0.05 <= t <= 0.95 and 0.05 <= u <= 0.95:
        return (a1[0] + t * d1x, a1[1] + t * d1y)
    return None

def detect_openings(msp, box):
    """الفتحات = مربّعات X على طبقة 'S - dot': كل فتحة = تقاطع قُطرَيها (تنفصل المتلاصقة). تشمل السلالم."""
    segs = []
    for e in msp.query("LINE LWPOLYLINE"):
        if e.dxf.layer != "S - dot": continue
        try:
            pts = [(e.dxf.start.x, e.dxf.start.y), (e.dxf.end.x, e.dxf.end.y)] if e.dxftype() == "LINE" \
                  else [(p[0], p[1]) for p in e.get_points("xy")]
        except Exception: continue
        for i in range(len(pts) - 1):
            a, b = pts[i], pts[i + 1]
            if in_box(a[0], a[1], box) or in_box(b[0], b[1], box): segs.append((a, b))
    # كل فتحة = تقاطع قُطرين متعاكسي الميل. نكتشف التقاطعات (لا مكوّنات متصلة) فتنفصل الفتحات المتلاصقة.
    diags = [s for s in segs if abs(s[0][0] - s[1][0]) > 400 and abs(s[0][1] - s[1][1]) > 400]
    raw = []
    for i in range(len(diags)):
        a1, b1 = diags[i]; s1 = 1 if (b1[0] - a1[0]) * (b1[1] - a1[1]) > 0 else -1
        for j in range(i + 1, len(diags)):
            a2, b2 = diags[j]; s2 = 1 if (b2[0] - a2[0]) * (b2[1] - a2[1]) > 0 else -1
            if s1 == s2: continue
            X = _seg_cross(a1, b1, a2, b2)
            if not X: continue
            xs = [a1[0], b1[0], a2[0], b2[0]]; ys = [a1[1], b1[1], a2[1], b2[1]]
            w = max(xs) - min(xs); h = max(ys) - min(ys)
            if w < 500 or h < 500: continue
            raw.append((w * h / 1e6, X[0], X[1], w, h))
    raw.sort(reverse=True)
    kept = []
    for a2, cx, cy, w, h in raw:   # نفس الـX قد يظهر بأكثر من زوج قُطري → ادمج المتطابق المركز فقط (<0.8م)
        if any(abs(cx - k[1]) < 800 and abs(cy - k[2]) < 800 for k in kept): continue
        kept.append((a2, cx, cy, w, h))
    total = sum(k[0] for k in kept)
    items = [(round(a2, 1), round(cx), round(cy), round(w), round(h)) for a2, cx, cy, w, h in kept]
    return total, items

def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("ceiling")
    ap.add_argument("--dxf", default=takeoff.DEFAULT_DXF)
    ap.add_argument("--outdir", default=takeoff.DEFAULT_OUT)
    ap.add_argument("--gx", type=float, help="العرض الكلّي (م)")
    ap.add_argument("--gy", type=float, help="الطول الكلّي (م)")
    ap.add_argument("--area", type=float, help="المساحة الإجمالية (م²) بدل الأبعاد")
    ap.add_argument("--openings", type=float, help="إجمالي مساحة الفتحات (م²) — تحدّدها أنت")
    ap.add_argument("--lap", type=int, default=DEF_LAP_MULT, help="مضاعف طول وصلة التراكب (×القطر). 60 شد / 40 ضغط")
    ap.add_argument("--conservative", action="store_true", help="احتسب قصاصة آخر سيخ لكل مجرى (بلا إعادة استخدام). الافتراضي: تراكب فقط")
    args = ap.parse_args()
    args.reuse = not args.conservative   # الافتراضي = محسّن (تراكب فقط، إعادة استخدام القصاصة)

    doc = ezdxf.readfile(args.dxf); msp = doc.modelspace()
    titles = takeoff.all_titles(msp)
    hit = takeoff.find_title(titles, args.ceiling, "TOP")
    if not hit:
        print("لم أجد لوحة TOP لهذا السقف."); return
    tx, ty, ttl = hit[0]
    box = takeoff.plan_bbox(titles, tx, ty)

    notes = region_notes(msp, box)
    # الشبكة السفلية قد تُوضع على لوحة الإضافي السفلي؛ نبحثها هناك أيضاً لضمان التقاط الطبقتين
    bhit = takeoff.find_title(titles, args.ceiling, "BOTTOM")
    if bhit:
        bx, by, _ = bhit[0]
        notes += region_notes(msp, takeoff.plan_bbox(titles, bx, by))
    # مواصفات مميّزة لكل (وجه، قطر، تباعد) — نجمع الطبقتين ونزيل التكرار
    seen = set(); dn = []
    for n in notes:
        k = (n["face"], n["dia"], n["sp"])
        if k in seen: continue
        seen.add(k); dn.append(n)
    notes = dn
    if not notes:
        print("لم أجد ملاحظات TOP=/BOT.= (الفرش/الغطاء) في هذه اللوحة."); return
    faces = {n["face"] for n in notes}
    if faces != {"TOP", "BOT"}:
        print(f"تنبيه: وُجد وجه واحد فقط {faces} — تأكّد أن الطبقتين (علوية+سفلية) ممثّلتان.")
    if args.area:
        gross = args.area; gx = gy = None
    elif args.gx and args.gy:
        gx, gy = args.gx, args.gy; gross = gx * gy
    else:
        gx, gy = rebar_extent(msp, box)   # قياس تلقائي من امتداد الحديد
        gross = gx * gy
    if args.openings is not None:
        openings, op_items, op_note = args.openings, [], " (يدوي)"
    else:
        openings, op_items = detect_openings(msp, box)   # مربّعات X على S-dot (شاملة السلالم)
        op_note = f" (تلقائي: {len(op_items)} فتحة X على S-dot، شاملة السلالم)"
    net = gross - openings

    # group steel length by diameter
    per_dia = collections.defaultdict(float)   # dia -> total length (m)
    detail = []
    for n in notes:
        lpm2 = 2000.0 / n["sp"]                 # two-way mesh, one layer
        length = net * lpm2
        per_dia[n["dia"]] += length
        detail.append((n["face"], n["dia"], n["sp"], round(lpm2, 2), round(length)))

    dim_str = f"{gx:.2f}×{gy:.2f}" if gx else f"مساحة مُدخَلة"
    print(f"اللوحة: {ttl}")
    print(f"أبعاد كلّية: {dim_str}  إجمالي={gross:.1f} م²")
    print(f"فتحات (شاملة السلالم)={openings:.1f}{op_note}  صافي={net:.1f} م²")
    for a2, cx, cy, *rest in op_items[:12]:
        wh = f"  {rest[0]}×{rest[1]}مم" if rest else ""
        print(f"     فتحة ~{a2} م² عند ({cx},{cy}){wh}")
    print("المواصفات:", [f"{n['face']} T{n['dia']}@{n['sp']}" for n in notes])
    lap_mode = "تراكب فقط (قصاصة مُعادة)" if args.reuse else "قياسي (يشمل قصاصة آخر سيخ)"
    print(f"\n=== الحصر (فرش+غطاء) === | وصلة {args.lap}×القطر — {lap_mode}")
    rows = []
    grand_bars = 0
    for dia in sorted(per_dia):
        length = per_dia[dia]                       # صافي طول الشبكة (الفتحات مخصومة)
        f = lap_factor(gx, gy, dia, args.lap, args.reuse)   # زيادة الوصلات (حسب القطر والأبعاد)
        length_w = length * f
        bars = math.ceil(length_w / STOCK)
        grand_bars += bars
        rows.append((dia, round(length), round(length_w), bars, f))
        print(f"  Ø{dia}: صافي={round(length)} م ×{f:.3f} (وصلات) = {round(length_w)} م → {bars} سيخ (12م)")

    # Excel
    NAVY, BLUE, LT = "1F3864", "2E5496", "D9E1F2"
    thin = Side(style="thin", color="BFBFBF"); bd = Border(thin, thin, thin, thin)
    wb = Workbook(); ws = wb.active; ws.title = "الفرش والغطاء"; ws.sheet_view.rightToLeft = True
    def st(c, *, b=False, sz=11, color="000000", fill=None):
        c.font = Font(bold=b, size=sz, color=color); c.alignment = Alignment("center", "center")
        if fill: c.fill = PatternFill("solid", fgColor=fill)
        c.border = bd
    ws.merge_cells("A1:E1"); st(ws["A1"], b=True, sz=14, color="FFFFFF", fill=NAVY)
    ws["A1"] = f"حصر الفرش والغطاء — {ttl}"
    ws.merge_cells("A2:E2"); st(ws["A2"], sz=9, color="FFFFFF", fill=BLUE)
    ws["A2"] = (f"إجمالي {dim_str}={gross:.0f}م² − فتحات شاملة السلالم {openings:.0f} = صافي {net:.0f}م² "
                f"| وصلات تراكب {args.lap}×القطر {'(تراكب فقط)' if args.reuse else '(قياسي)'} ÷12")
    hdr = ["القطر", "صافي الطول (م)", "مع الوصلات (م)", "عدد أسياخ 12م", "المواصفة"]
    for j, h in enumerate(hdr, 1): st(ws.cell(4, j, h), b=True, fill=NAVY, color="FFFFFF")
    r = 5
    specs_by_dia = collections.defaultdict(list)
    for n in notes: specs_by_dia[n["dia"]].append(f"{n['face']} T{n['dia']}@{n['sp']}")
    for dia, length, length_w, bars, f in rows:
        st(ws.cell(r, 1, f"Ø{dia}"), b=True, fill=LT); st(ws.cell(r, 2, length)); st(ws.cell(r, 3, length_w))
        st(ws.cell(r, 4, bars), b=True); st(ws.cell(r, 5, " ، ".join(sorted(set(specs_by_dia[dia])))))
        r += 1
    ws.column_dimensions["A"].width = 10
    for cc in "BCDE": ws.column_dimensions[cc].width = 18
    os.makedirs(args.outdir, exist_ok=True)
    out = os.path.join(args.outdir, takeoff.safe(ttl) + " - MESH.xlsx"); wb.save(out)
    print("\nSAVED:", out, "| إجمالي الأسياخ:", grand_bars)

if __name__ == "__main__":
    main()
