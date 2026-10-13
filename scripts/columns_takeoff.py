# -*- coding: utf-8 -*-
"""
حصر أعداد الأعمدة من لوحة COLUMNS & AXES في ملف الـ DXF.
Column COUNT takeoff from the COLUMNS & AXES plan.

يقرأ علامات الأعمدة (C, C1..C10) والكورات (CORE 1..n) من مساحة اللوحة،
يعدّها حسب الماركة، ويربط كل ماركة بمقطعها وتسليحها من جدول الأعمدة (COLUMN SCHEDULE)،
ثم يخرّج جدول Excel منظّم.

الاستخدام:
    python columns_takeoff.py
    python columns_takeoff.py --dxf "<path>" --out "<xlsx>"
"""
import argparse, os, re, math, collections, glob
import ezdxf
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# مسارات محمولة نسبةً للمستودع (scripts/ تحت الجذر، واللوحة في drawings/)
_ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
def _default_dxf():
    for p in sorted(glob.glob(os.path.join(_ROOT, "drawings", "*.dxf"))):
        return p
    return os.path.join(_ROOT, "drawings", "STR-MOTTAGAH.dxf")
DEFAULT_DXF = _default_dxf()
DEFAULT_OUT = os.path.join(_ROOT, "output", "حصر الأعمدة - COLUMNS & AXES.xlsx")

# علامة الماركة: C وحدها أو C متبوعة بأرقام (C1..C10)
MARK_RE = re.compile(r"^C\s*\.?\s*(\d{0,3})$", re.I)
CORE_RE = re.compile(r"^CORE\s*\.?\s*(\d+)$", re.I)
SIZE_RE = re.compile(r"SIZE\s*=\s*(\d+)\s*[xX]\s*(\d+)", re.I)
VERT_RE = re.compile(r"VERTICAL\s*=\s*(.+)$", re.I)

def text_of(e):
    try:
        t = e.plain_text() if e.dxftype() == "MTEXT" else e.dxf.text
    except Exception:
        t = getattr(e.dxf, "text", "")
    return " ".join((t or "").split())

def mark_key(m):
    """ترتيب منطقي: C ثم C1..C10 ثم الكورات."""
    if m == "C":
        return (0, 0)
    mm = re.match(r"C(\d+)$", m)
    if mm:
        return (0, int(mm.group(1)))
    cm = re.match(r"CORE(\d+)$", m)
    if cm:
        return (1, int(cm.group(1)))
    return (2, 0)

def collect(msp):
    """يفصل نصوص اللوحة (x<-480000) عن نصوص الجدول (x>-460000)."""
    plan, sched = [], []
    for e in msp.query("TEXT MTEXT"):
        try:
            ip = e.dxf.insert
            x, y = float(ip.x), float(ip.y)
        except Exception:
            continue
        t = text_of(e)
        if not t:
            continue
        if x < -480000 and 370000 < y < 430000:
            plan.append((x, y, t))
        elif -460000 < x < -330000 and 440000 < y < 490000:
            sched.append((x, y, t))
    return plan, sched

def count_marks(plan):
    counts = collections.Counter()
    for (x, y, t) in plan:
        mm = MARK_RE.match(t)
        cm = CORE_RE.match(t)
        if mm:
            counts["C" + mm.group(1)] += 1      # "C" أو "C3"
        elif cm:
            counts["CORE" + cm.group(1)] += 1
    return counts

def parse_schedule(sched):
    """لكل ماركة في الجدول: مقطع المرحلة الأولى (الأساس→سقف البدروم العلوي) + تسليح رأسي."""
    heads = []
    for (x, y, t) in sched:
        m = re.match(r"^C\s*(\d{0,3})$", t, re.I)
        if m and x < -419000:                    # عمود الماركات على أقصى يسار الجدول
            heads.append((x, y, "C" + m.group(1)))
    info = {}
    for (hx, hy, mark) in heads:
        # نصوص نفس الصف: العنوان أعلى الخلية والمقطع/التسليح أسفله بنحو 1000
        row = [(x, y, t) for (x, y, t) in sched if hy - 1450 < y < hy + 250]
        sizes = sorted([(x, SIZE_RE.search(t)) for (x, y, t) in row if SIZE_RE.search(t)])
        verts = sorted([(x, VERT_RE.search(t)) for (x, y, t) in row if VERT_RE.search(t)])
        size = f"{sizes[0][1].group(1)}×{sizes[0][1].group(2)}" if sizes else "—"
        vert = verts[0][1].group(1).strip() if verts else "—"
        info[mark] = (size, vert)
    return info

# ------------------------- Excel -------------------------
def build_excel(counts, info, path):
    marks = sorted([m for m in counts if not m.startswith("CORE")], key=mark_key)
    cores = sorted([m for m in counts if m.startswith("CORE")], key=mark_key)

    wb = Workbook(); ws = wb.active; ws.title = "حصر الأعمدة"; ws.sheet_view.rightToLeft = True
    NAVY, BLUE, LT, GREY = "1F3864", "2E5496", "D9E1F2", "F2F2F2"
    thin = Side(style="thin", color="BFBFBF"); bd = Border(thin, thin, thin, thin)
    def st(c, *, b=False, sz=11, color="000000", fill=None, ha="center", wrap=False):
        c.font = Font(bold=b, size=sz, color=color, name="Calibri")
        c.alignment = Alignment(horizontal=ha, vertical="center", wrap_text=wrap)
        if fill: c.fill = PatternFill("solid", fgColor=fill)
        c.border = bd

    ws.merge_cells("A1:D1"); st(ws["A1"], b=True, sz=14, color="FFFFFF", fill=NAVY)
    ws["A1"] = "حصر الأعمدة — لوحة COLUMNS & AXES"
    ws.merge_cells("A2:D2"); st(ws["A2"], sz=9, color="FFFFFF", fill=BLUE)
    ws["A2"] = "عدد الأعمدة حسب الماركة | المقطع والتسليح من جدول الأعمدة (المرحلة: الأساس ← سقف البدروم العلوي)"
    ws.row_dimensions[1].height = 26

    hdr = ["ماركة العمود", "العدد", "المقطع (مم)", "التسليح الرأسي"]
    for j, h in enumerate(hdr, 1):
        st(ws.cell(4, j, h), b=True, fill=NAVY, color="FFFFFF")

    r = 5; total = 0
    for m in marks:
        size, vert = info.get(m, ("—", "—"))
        st(ws.cell(r, 1, m), b=True, fill=GREY)
        st(ws.cell(r, 2, counts[m]), b=True)
        st(ws.cell(r, 3, size))
        st(ws.cell(r, 4, vert))
        total += counts[m]; r += 1
    st(ws.cell(r, 1, "إجمالي الأعمدة"), b=True, fill=NAVY, color="FFFFFF")
    st(ws.cell(r, 2, total), b=True, fill=LT)
    st(ws.cell(r, 3, "")); st(ws.cell(r, 4, "")); r += 2

    if cores:
        ws.merge_cells(f"A{r}:D{r}"); st(ws.cell(r, 1, "الكورات (حوائط قص) — Cores"), b=True, fill=BLUE, color="FFFFFF"); r += 1
        st(ws.cell(r, 1, "الماركة"), b=True, fill=GREY); st(ws.cell(r, 2, "العدد"), b=True, fill=GREY)
        st(ws.cell(r, 3, "")); st(ws.cell(r, 4, "")); r += 1
        ctot = 0
        for m in cores:
            lbl = m.replace("CORE", "CORE ")
            st(ws.cell(r, 1, lbl)); st(ws.cell(r, 2, counts[m])); st(ws.cell(r, 3, "")); st(ws.cell(r, 4, "")); ctot += counts[m]; r += 1
        st(ws.cell(r, 1, "إجمالي الكورات"), b=True, fill=NAVY, color="FFFFFF"); st(ws.cell(r, 2, ctot), b=True, fill=LT)
        st(ws.cell(r, 3, "")); st(ws.cell(r, 4, "")); r += 2

    ws.merge_cells(f"A{r}:D{r}")
    st(ws.cell(r, 1, "ملاحظة: «C» ماركة مستقلة (300×500، 8T16) غير «C1» (300×600، 10T16). المقاطع تتدرّج لأعلى بالأدوار — راجع جدول الأعمدة للمراحل."),
       sz=9, color="1F3864", fill="E2EFDA", ha="right", wrap=True); ws.row_dimensions[r].height = 34

    ws.column_dimensions["A"].width = 18
    ws.column_dimensions["B"].width = 10
    ws.column_dimensions["C"].width = 16
    ws.column_dimensions["D"].width = 18
    os.makedirs(os.path.dirname(path), exist_ok=True)
    wb.save(path)

def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--dxf", default=DEFAULT_DXF)
    ap.add_argument("--out", default=DEFAULT_OUT)
    args = ap.parse_args()

    doc = ezdxf.readfile(args.dxf); msp = doc.modelspace()
    plan, sched = collect(msp)
    counts = count_marks(plan)
    info = parse_schedule(sched)

    marks = sorted([m for m in counts if not m.startswith("CORE")], key=mark_key)
    cores = sorted([m for m in counts if m.startswith("CORE")], key=mark_key)
    print("=== عدد الأعمدة ===")
    tot = 0
    for m in marks:
        size, vert = info.get(m, ("—", "—"))
        print(f"  {m:5s} : {counts[m]:3d}   [{size}  {vert}]"); tot += counts[m]
    print(f"  إجمالي الأعمدة: {tot}")
    if cores:
        print("=== الكورات ===")
        for m in cores:
            print(f"  {m.replace('CORE','CORE ')} : {counts[m]}")

    build_excel(counts, info, args.out)
    print("\nSAVED:", args.out)

if __name__ == "__main__":
    main()
