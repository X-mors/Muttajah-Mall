# -*- coding: utf-8 -*-
"""
حصر خرسانة الأعمدة (Column CONCRETE takeoff) — الوضع: الأعمدة.
يبني على columns_takeoff: العدد لكل ماركة + مقطع كل مرحلة من المراحل الخمس من جدول الأعمدة،
ثم الحجم = Σ لكل (ماركة، مرحلة)  العدد × مساحة المقطع × ارتفاع المرحلة الصافي.

قاعدة المهندس للارتفاع: من سقف الدور الأسفل حتى سقف الدور (سقف↔سقف)،
والارتفاع الصافي للعمود = ارتفاع المرحلة − سُمك ما يعلو العمود (بلاطة ٢٢سم أو drop panel ٣٢سم).
المناسيب/الارتفاعات وسُمك البلاطة والـ drop panel تُدخَل من المهندس (غير موجودة نصّاً باللوحة).

الاستخدام:
    python concrete_columns_takeoff.py --debug          # يعرض مقطع كل مرحلة لكل ماركة (تحقّق)
    python concrete_columns_takeoff.py --heights 3.95,3.20,6.40,6.40,3.20 --slab 0.22
"""
import argparse, os, re, math, collections
import ezdxf
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import columns_takeoff as ct   # إعادة استخدام: collect, count_marks, mark_key, text_of, SIZE_RE

DEFAULT_DXF = ct.DEFAULT_DXF
DEFAULT_OUT = r"D:\مشروع الحصر\output\حصر خرسانة الأعمدة - COLUMNS & AXES.xlsx"

STAGE_HDR_RE = re.compile(r"^FROM\s+.+\b(CEILING|ROOF|FOUNDATION)", re.I)

def stage_anchors(sched):
    """مراكز X للمراحل الخمس من عناوينها 'FROM ... TO ... CEILING' (مرتّبة يسار→يمين = مرحلة 1→5)."""
    hs = []
    for (x, y, t) in sched:
        tt = " ".join(t.split())
        if STAGE_HDR_RE.match(tt) and "SCHEDULE" not in tt.upper():
            hs.append((x, tt))
    # دمج التكرار بتقريب X، ترتيب تصاعدي (الأكثر سلبيةً = المرحلة الأولى)
    uniq = {}
    for x, t in hs:
        uniq[round(x / 500) * 500] = (x, t)
    ordered = sorted(uniq.values(), key=lambda p: p[0])
    return ordered   # [(x, name)] من المرحلة 1 إلى 5

def parse_schedule_stages(sched):
    """لكل ماركة: قائمة بطول 5 = مقطع كل مرحلة (W,H) بالمم أو None إذا منتهٍ/غير موجود (----)."""
    anchors = stage_anchors(sched)
    ax = [a[0] for a in anchors]
    names = [a[1] for a in anchors]
    # رؤوس الماركات في عمود الماركات أقصى اليسار
    heads = []
    for (x, y, t) in sched:
        m = re.match(r"^C\s*\.?\s*(\d{0,3})$", t, re.I)
        if m and x < min(ax) - 300:               # يسار أول مرحلة
            heads.append((x, y, "C" + m.group(1)))
    out = {}
    for (hx, hy, mark) in heads:
        row = [(x, t) for (x, yy, t) in sched if hy - 1450 < yy < hy + 250]
        stages = [None] * len(ax)
        for (x, t) in row:
            ms = ct.SIZE_RE.search(t)
            if not ms:
                continue
            # أقرب مرحلة بالـ X
            k = min(range(len(ax)), key=lambda i: abs(ax[i] - x))
            w, h = int(ms.group(1)), int(ms.group(2))
            stages[k] = (w, h)
        out[mark] = stages
    return out, names

def section_area_m2(wh):
    """مساحة المقطع م². wh=(W_mm,H_mm). دائري لو W==H==D؟ نُبقيه مستطيلاً افتراضاً."""
    if not wh:
        return 0.0
    w, h = wh
    return (w / 1000.0) * (h / 1000.0)

def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--dxf", default=DEFAULT_DXF)
    ap.add_argument("--out", default=DEFAULT_OUT)
    ap.add_argument("--heights", help="ارتفاعات المراحل الخمس سقف↔سقف (م)، مثال: 3.95,3.20,6.40,6.40,3.20")
    ap.add_argument("--slab", type=float, default=0.22, help="سُمك البلاطة العلوية (م) يُخصم من الارتفاع (افتراضي 0.22)")
    ap.add_argument("--drop", type=float, default=0.32, help="سُمك الـ drop panel (م) حيث يوجد (افتراضي 0.32)")
    ap.add_argument("--drop-marks", default="", help="ماركات تحتها drop panel (تُخصم 0.32 بدل البلاطة)، مثال: C1,C6")
    ap.add_argument("--mode", choices=["gross", "net"], default="net", help="gross=سقف↔سقف كاملاً، net=بخصم البلاطة/الهابط")
    ap.add_argument("--waste", type=float, default=5.0, help="نسبة الهالك للخرسانة ٪ (افتراضي 5)")
    ap.add_argument("--debug", action="store_true", help="اعرض مقطع كل مرحلة لكل ماركة فقط")
    args = ap.parse_args()

    doc = ezdxf.readfile(args.dxf); msp = doc.modelspace()
    plan, sched = ct.collect(msp)
    counts = ct.count_marks(plan)
    stages_by_mark, stage_names = parse_schedule_stages(sched)

    marks = sorted([m for m in counts if not m.startswith("CORE")], key=ct.mark_key)

    # عرض مقاطع المراحل (تحقّق) — يعمل دائماً
    print("=== مقطع كل مرحلة لكل ماركة (مم) ===")
    print("المرحلة:", " | ".join(f"{i+1}" for i in range(len(stage_names))))
    for m in marks:
        st = stages_by_mark.get(m, [])
        cells = []
        for s in st:
            cells.append(f"{s[0]}×{s[1]}" if s else "----")
        print(f"  {m:5s} (عدد {counts[m]:2d}): " + " | ".join(f"{c:>9s}" for c in cells))
    print("\nأسماء المراحل:")
    for i, n in enumerate(stage_names, 1):
        print(f"  {i}. {n}")

    if args.debug or not args.heights:
        print("\n(للحساب: مرّر --heights بخمس قيم سقف↔سقف بالمتر، و--slab سُمك البلاطة.)")
        return

    heights = [float(x) for x in args.heights.split(",")]
    if len(heights) != len(stage_names):
        print(f"تنبيه: عدد الارتفاعات ({len(heights)}) لا يساوي عدد المراحل ({len(stage_names)})."); return
    drop_marks = set(m.strip().upper() for m in args.drop_marks.split(",") if m.strip())

    # الحساب
    rows = []           # (stage_idx, mark, count, W, H, area, h_storey, deduct, h_clear, vol)
    stage_tot = collections.defaultdict(float)
    grand = 0.0
    for si in range(len(stage_names)):
        for m in marks:
            wh = stages_by_mark.get(m, [None] * len(stage_names))[si]
            if not wh:
                continue
            n = counts[m]
            area = section_area_m2(wh)
            hS = heights[si]
            ded = (args.drop if m.upper() in drop_marks else args.slab) if args.mode == "net" else 0.0
            hC = max(0.0, hS - ded)
            vol = n * area * hC
            rows.append((si, m, n, wh[0], wh[1], area, hS, ded, hC, vol))
            stage_tot[si] += vol; grand += vol

    print("\n=== حصر خرسانة الأعمدة ===")
    for si in range(len(stage_names)):
        if stage_tot[si] == 0: continue
        print(f"\n▶ المرحلة {si+1}: {stage_names[si]}  (ارتفاع {heights[si]}م)")
        for (s, m, n, W, H, area, hS, ded, hC, vol) in [r for r in rows if r[0] == si]:
            print(f"   {m:5s}: {n:2d} × {W}×{H} ({area:.3f}م²) × {hC:.2f}م = {vol:.2f} م³")
        print(f"   ── إجمالي المرحلة: {stage_tot[si]:.2f} م³")
    wst = grand * (1 + args.waste / 100.0)
    print(f"\n الإجمالي ({args.mode}): {grand:.2f} م³  |  بعد الهالك +{args.waste:.0f}٪: {wst:.2f} م³")

    build_excel(rows, stage_tot, stage_names, heights, counts, args, grand, wst)
    print("SAVED:", args.out)

def build_excel(rows, stage_tot, stage_names, heights, counts, args, grand, wst):
    NAVY, BLUE, LT, GREY, GRN = "1F3864", "2E5496", "D9E1F2", "F2F2F2", "E2EFDA"
    thin = Side(style="thin", color="BFBFBF"); bd = Border(thin, thin, thin, thin)
    wb = Workbook(); ws = wb.active; ws.title = "حصر خرسانة الأعمدة"; ws.sheet_view.rightToLeft = True
    def st(c, *, b=False, sz=11, color="000000", fill=None, ha="center", wrap=False):
        c.font = Font(bold=b, size=sz, color=color); c.alignment = Alignment(ha, "center", wrap_text=wrap)
        if fill: c.fill = PatternFill("solid", fgColor=fill)
        c.border = bd
    ws.merge_cells("A1:G1"); st(ws["A1"], b=True, sz=14, color="FFFFFF", fill=NAVY)
    ws["A1"] = "حصر خرسانة الأعمدة — COLUMNS & AXES"
    ws.merge_cells("A2:G2"); st(ws["A2"], sz=9, color="FFFFFF", fill=BLUE)
    ws["A2"] = f"الطريقة: {'صافي (بخصم البلاطة/الهابط)' if args.mode=='net' else 'إجمالي سقف↔سقف'} | هالك +{args.waste:.0f}٪ | الارتفاعات بإدخال المهندس"
    hdr = ["الماركة", "العدد", "المقطع (مم)", "المساحة (م²)", "الارتفاع الصافي (م)", "الحجم (م³)", "المرحلة"]
    r = 4
    for si in range(len(stage_names)):
        srows = [x for x in rows if x[0] == si]
        if not srows: continue
        ws.merge_cells(f"A{r}:G{r}"); st(ws.cell(r, 1, f"◄ المرحلة {si+1}: {stage_names[si]}  (ارتفاع {heights[si]}م)"), b=True, color="FFFFFF", fill=BLUE, ha="right"); r += 1
        for j, h in enumerate(hdr, 1): st(ws.cell(r, j, h), b=True, fill=NAVY, color="FFFFFF")
        r += 1
        for (s, m, n, W, H, area, hS, ded, hC, vol) in srows:
            st(ws.cell(r, 1, m), b=True, fill=GREY); st(ws.cell(r, 2, n)); st(ws.cell(r, 3, f"{W}×{H}"))
            st(ws.cell(r, 4, round(area, 3))); st(ws.cell(r, 5, round(hC, 2))); st(ws.cell(r, 6, round(vol, 2)))
            st(ws.cell(r, 7, f"{si+1}")); r += 1
        st(ws.cell(r, 1, "إجمالي المرحلة"), b=True, fill=LT); st(ws.cell(r, 6, round(stage_tot[si], 2)), b=True, fill=LT)
        for cc in (2, 3, 4, 5, 7): st(ws.cell(r, cc, None), fill=LT)
        r += 2
    st(ws.cell(r, 1, "الإجمالي الكلي (م³)"), b=True, fill=NAVY, color="FFFFFF"); st(ws.cell(r, 6, round(grand, 2)), b=True, fill=LT); r += 1
    st(ws.cell(r, 1, f"بعد الهالك +{args.waste:.0f}٪"), b=True, fill=GREY); st(ws.cell(r, 6, round(wst, 2)), b=True, fill=GREY); r += 2
    ws.merge_cells(f"A{r}:G{r}")
    st(ws.cell(r, 1, "القاعدة: ارتفاع العمود من سقف الدور الأسفل إلى سقف الدور (سقف↔سقف) ناقص سُمك البلاطة (٢٢سم) أو الـ drop panel (٣٢سم) حيث يوجد. المقطع من جدول الأعمدة لكل مرحلة. الكورات تُحصر منفصلة."), sz=9, fill=GRN, ha="right", wrap=True); ws.row_dimensions[r].height = 40
    ws.column_dimensions["A"].width = 12
    for cc in "BCDEFG": ws.column_dimensions[cc].width = 16
    os.makedirs(os.path.dirname(args.out), exist_ok=True); wb.save(args.out)

if __name__ == "__main__":
    main()
