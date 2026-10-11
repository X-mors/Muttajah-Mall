# -*- coding: utf-8 -*-
"""
حصر سقف كامل: التسليح الإضافي العلوي + السفلي في شيت إكسل واحد منسّق،
وملف مراجعة DXF/DWG واحد عليه دوائر مرقّمة (كل دائرة رقمها = صفّها في «التفاصيل»).

الاستخدام:
    python ceiling_takeoff.py "LOWER BASEMENT CEILING"
    python ceiling_takeoff.py "LOWER BASEMENT CEILING" --dwg
"""
import argparse, os, math, collections
import ezdxf
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import takeoff  # المحرّك الأساسي (نفس المجلد)

def load_plan(msp, titles, ceiling, extra):
    hits = takeoff.find_title(titles, ceiling, extra)
    if not hits:
        return None
    tx, ty, ttl = hits[0]
    box = takeoff.plan_bbox(titles, tx, ty)
    return dict(title=ttl, box=box, callouts=takeoff.extract(msp, box))

def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("ceiling", help="اسم السقف، مثل: LOWER BASEMENT CEILING")
    ap.add_argument("--dxf", default=takeoff.DEFAULT_DXF)
    ap.add_argument("--outdir", default=takeoff.DEFAULT_OUT)
    ap.add_argument("--dwg", action="store_true")
    args = ap.parse_args()

    doc = ezdxf.readfile(args.dxf); msp = doc.modelspace()
    titles = takeoff.all_titles(msp)
    top = load_plan(msp, titles, args.ceiling, "TOP")
    bot = load_plan(msp, titles, args.ceiling, "BOTTOM")
    if not top or not bot:
        print("تعذّر إيجاد لوحتي العلوي/السفلي لهذا السقف."); return

    # ترقيم تسلسلي مستقر: العلوي أولاً ثم السفلي، مرتّبين مكانياً (أعلى→أسفل، يسار→يمين)
    serial = 0
    sections = [("علوي", "TOP", top), ("سفلي", "BOTTOM", bot)]
    for _, _, pl in sections:
        for c in sorted(pl["callouts"], key=lambda c: (-c["y"], c["x"])):
            serial += 1; c["serial"] = serial

    # ---------- إكسل ----------
    NAVY, BLUE, LT, GREY, AMB, GRN = "1F3864", "2E5496", "D9E1F2", "F2F2F2", "FFF2CC", "E2EFDA"
    thin = Side(style="thin", color="BFBFBF"); bd = Border(thin, thin, thin, thin)
    wb = Workbook()
    def st(c, *, b=False, sz=11, color="000000", fill=None, ha="center", wrap=False):
        c.font = Font(bold=b, size=sz, color=color)
        c.alignment = Alignment(horizontal=ha, vertical="center", wrap_text=wrap)
        if fill: c.fill = PatternFill("solid", fgColor=fill)
        c.border = bd

    ws = wb.active; ws.title = "الملخص"; ws.sheet_view.rightToLeft = True
    ws.merge_cells("A1:F1"); st(ws["A1"], b=True, sz=15, color="FFFFFF", fill=NAVY)
    ws["A1"] = f"حصر أعداد أسياخ التسليح الإضافي — {args.ceiling}"
    ws.merge_cells("A2:F2"); st(ws["A2"], sz=9, color="FFFFFF", fill=BLUE)
    ws["A2"] = "العلوي + السفلي | العدد = عرض التوزيع ÷ التباعد (تقريب لأعلى) | السلالم مستبعدة"
    ws.row_dimensions[1].height = 26
    grand = collections.Counter()
    r = 4
    for label, _, pl in sections:
        slab = [c for c in pl["callouts"] if not c["stair"]]
        stair_n = sum(1 for c in pl["callouts"] if c["stair"])
        dias = sorted({c["dia"] for c in slab})
        piv = collections.defaultdict(lambda: {d: 0 for d in dias}); tot = {d: 0 for d in dias}
        for c in slab:
            if not c["count"]: continue
            Lk = c["L"] if c["L"] else "unknown"
            piv[Lk][c["dia"]] += c["count"]; tot[c["dia"]] += c["count"]; grand[c["dia"]] += c["count"]
        lengths = sorted([k for k in piv if k != "unknown"]) + (["unknown"] if "unknown" in piv else [])
        last = get_column_letter(1 + max(len(dias), 1))
        ws.merge_cells(f"A{r}:{last}{r}"); st(ws.cell(r, 1, f"◄ اللوحة: التسليح الإضافي {label}"), b=True, color="FFFFFF", fill=BLUE, ha="right"); r += 1
        st(ws.cell(r, 1, "طول القطعية (م)"), b=True, fill=NAVY, color="FFFFFF")
        for j, d in enumerate(dias): st(ws.cell(r, 2 + j, f"Ø{d}"), b=True, fill=NAVY, color="FFFFFF")
        r += 1
        for Lk in lengths:
            lbl = "غير محدّد ⚠" if Lk == "unknown" else f"{Lk/1000:g}"
            st(ws.cell(r, 1, lbl), fill=(AMB if Lk == "unknown" else GREY))
            for j, d in enumerate(dias):
                v = piv[Lk][d]; st(ws.cell(r, 2 + j, v if v else None), fill=(AMB if Lk == "unknown" else None))
            r += 1
        st(ws.cell(r, 1, "الإجمالي"), b=True, fill=LT)
        for j, d in enumerate(dias): st(ws.cell(r, 2 + j, tot[d] or None), b=True, fill=LT)
        r += 1
        red = sum(1 for c in slab if c["count"] is None); cyan = sum(1 for c in slab if c["count"] and c.get("src") == "heur")
        ws.merge_cells(f"A{r}:{last}{r}"); st(ws.cell(r, 1, f"سلالم مستبعدة: {stair_n} | بحاجة مراجعة: أحمر {red}، عرض غير مؤكّد {cyan}"), sz=9, fill=GRN, ha="right"); r += 2
    # grand total by diameter
    gd = sorted(grand)
    st(ws.cell(r, 1, "الإجمالي الكلي (علوي+سفلي)"), b=True, color="FFFFFF", fill=NAVY)
    for j, d in enumerate(gd): st(ws.cell(r, 2 + j, f"Ø{d}: {grand[d]}"), b=True, fill=LT)
    ws.column_dimensions["A"].width = 18
    for j in range(1, 8): ws.column_dimensions[get_column_letter(1 + j)].width = 12

    # ---------- ورقة التفاصيل (الرقم يطابق الدائرة) ----------
    ws2 = wb.create_sheet("التفاصيل"); ws2.sheet_view.rightToLeft = True
    hdr = ["الرقم", "اللوحة", "القطر", "التباعد", "الوجه", "الطول L (مم)", "عرض التوزيع (مم)", "العدد", "النوع", "x", "y"]
    for j, h in enumerate(hdr, 1): st(ws2.cell(1, j, h), b=True, fill=NAVY, color="FFFFFF")
    allc = []
    for label, _, pl in sections:
        for c in pl["callouts"]:
            c["_plan"] = label; allc.append(c)
    i = 1
    for c in sorted(allc, key=lambda c: c["serial"]):
        typ = "سلّم (مستبعد)" if c["stair"] else "بلاطة"
        fill = "E7E6E6" if c["stair"] else (None if c.get("src") in ("circle", "barline") else AMB)
        row = [c["serial"], c["_plan"], f"Ø{c['dia']}", c["sp"], ("علوي" if c["face"] == "T" else "سفلي"),
               c["L"], c["w"], (None if c["stair"] else c["count"]), typ, round(c["x"]), round(c["y"])]
        for j, v in enumerate(row, 1): st(ws2.cell(i + 1, j, v), fill=fill)
        i += 1
    for j, wd in enumerate([6, 7, 7, 8, 7, 12, 14, 7, 12, 10, 10], 1):
        ws2.column_dimensions[get_column_letter(j)].width = wd
    ws2.freeze_panes = "A2"

    os.makedirs(args.outdir, exist_ok=True)
    base = os.path.join(args.outdir, takeoff.safe(args.ceiling) + " - TOP+BOTTOM")
    xlsx = base + ".xlsx"; wb.save(xlsx)

    # ---------- ملف مراجعة واحد بالدوائر المرقّمة ----------
    doc2 = ezdxf.readfile(args.dxf); msp2 = doc2.modelspace()
    LAYER = "CHECK-COUNT"
    if LAYER not in doc2.layers: doc2.layers.add(LAYER, color=7)
    R = 850; GREEN, CYAN, RED, MAG = 3, 4, 1, 6
    for c in allc:
        if c["stair"]: col = MAG
        elif c["count"] is None: col = RED
        elif c.get("src") == "heur": col = CYAN
        else: col = GREEN
        msp2.add_circle((c["x"], c["y"]), R, dxfattribs={"layer": LAYER, "color": col})
        tt = msp2.add_text(str(c["serial"]), height=300, dxfattribs={"layer": LAYER, "color": col})
        tt.dxf.insert = (c["x"] + R * 0.7, c["y"] + R * 0.7)
    review = base + " - review.dxf"; doc2.saveas(review)
    print("SAVED:", xlsx); print("SAVED:", review)
    print("الإجمالي الكلي:", {f"Ø{d}": grand[d] for d in gd})
    if args.dwg:
        d = takeoff.to_dwg(review); print("DWG:", d or "فشل")

if __name__ == "__main__":
    main()
